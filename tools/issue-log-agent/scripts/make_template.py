"""Revisi Template Issue Backlog v1 menjadi v2 sesuai config agent.

Per tab, kolom disusun ulang menjadi:
    Title | Source | <sections dari config> | Issue Board Status | <kolom write-back GitLab>

- Style header (font, warna, tinggi baris) dan warna tab diambil dari v1.
- Kolom write-back diberi header abu-abu karena diisi agent.
- Dropdown untuk Source (per tipe) dan Issue Board Status.
- Data yang sudah ada di v1 dipindahkan berdasarkan nama header, jadi tidak ada isi yang hilang.

    pip install openpyxl
    python scripts/make_template.py --from templates/Template_Backlog_v1.xlsx \
        --out templates/Template_Backlog_v2.xlsx
"""

from __future__ import annotations

import argparse
import sys
from copy import copy
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from issue_log_agent.config import load_config  # noqa: E402

AGENT_FILL = PatternFill("solid", fgColor="FF666666")
MAX_ROWS = 1000
WIDTH_TITLE, WIDTH_SOURCE, WIDTH_SECTION, WIDTH_STATUS, WIDTH_AGENT = 45, 12, 27, 18, 18


def _norm(h: object) -> str:
    return str(h or "").strip().rstrip("*").strip().lower()


def revise(src: Path, out: Path, config: Path) -> None:
    cfg = load_config(config, require_token=False)
    wb = load_workbook(src)
    statuses = list(cfg.statuses)
    title_col, source_col, status_col = (cfg.common_columns[k] for k in ("title", "source", "status"))

    for t in cfg.types:
        if t.worksheet not in wb.sheetnames:
            wb.create_sheet(t.worksheet)
        ws = wb[t.worksheet]

        old_headers = [c.value for c in ws[1]]
        header_style = ws.cell(1, 1)._style if ws.max_column and ws.cell(1, 1).has_style else None
        header_height = ws.row_dimensions[1].height
        rows = [
            {_norm(h): v for h, v in zip(old_headers, r) if h}
            for r in ws.iter_rows(min_row=2, values_only=True)
            if any(v not in (None, "") for v in r)
        ]

        inputs = [title_col, source_col] + [s.column for s in t.sections] + [status_col]
        agent = list(cfg.writeback.values())
        headers = inputs + agent
        known = {_norm(h) for h in headers}
        lost = [h for h in old_headers if h and _norm(h) not in known]
        if lost:
            raise SystemExit(f"[{t.worksheet}] kolom v1 tidak ada di config, data bisa hilang: {lost}")

        ws.delete_cols(1, max(ws.max_column, len(headers)))
        ws.data_validations.dataValidation = []
        for i, name in enumerate(headers, 1):
            c = ws.cell(1, i, name)
            if header_style is not None:
                c._style = copy(header_style)
            if i > len(inputs):
                font = copy(c.font)
                font.color = "FFFFFFFF"
                c.font, c.fill = font, AGENT_FILL
            if name == title_col:
                width = WIDTH_TITLE
            elif name == source_col:
                width = WIDTH_SOURCE
            elif name == status_col:
                width = WIDTH_STATUS
            elif i > len(inputs):
                width = WIDTH_AGENT
            else:
                width = WIDTH_SECTION
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[1].height = header_height or 31.5

        for r, data in enumerate(rows, 2):
            for i, name in enumerate(headers, 1):
                ws.cell(r, i, data.get(_norm(name)))

        # Notes on headers so reporters know the rules without opening the guideline.
        required = {s.column for s in t.sections if s.required}
        notes = {
            title_col: "Judul tanpa prefix. Prefix [" + t.name + "][SOURCE]- ditambahkan otomatis.\n"
                       "Jika kosong dan fitur Claude aktif, agent membuatkan judul untuk dicek.",
            source_col: "Wajib. Pilih: " + ", ".join(t.sources),
            status_col: "Pilih: " + ", ".join(statuses) + ". Kosong = " + cfg.default_status + ".",
            **{col: "Wajib diisi." for col in required},
            **{col: "Diisi otomatis oleh agent. Jangan diubah." for col in agent},
        }
        for i, name in enumerate(headers, 1):
            if name in notes:
                ws.cell(1, i).comment = Comment(notes[name], "Issue Backlog Agent")

        def dropdown(col_name: str, values: list[str]) -> None:
            col = get_column_letter(headers.index(col_name) + 1)
            dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle=col_name,
                                error="Pilih salah satu: " + ", ".join(values))
            dv.add(f"{col}2:{col}{MAX_ROWS}")
            ws.add_data_validation(dv)

        dropdown(source_col, t.sources)
        dropdown(status_col, statuses)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Template v2 ditulis ke {out}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--from", dest="src", default="templates/Template_Backlog_v1.xlsx")
    ap.add_argument("--out", default="templates/Template_Backlog_v2.xlsx")
    ap.add_argument("--config", default="config.example.yaml")
    args = ap.parse_args()
    revise(Path(args.src), Path(args.out), Path(args.config))


if __name__ == "__main__":
    main()
